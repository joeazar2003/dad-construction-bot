import os
import re
from collections import defaultdict
from datetime import datetime

import requests
from flask import Flask, request
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

import github_store
from category import detect_category, category_list

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

app = Flask(__name__)

TELEGRAM_TOKEN = os.environ["TELEGRAM_TOKEN"]
# Comma-separated numeric Telegram chat IDs allowed to use this bot, e.g. "111111,222222"
# (your dad's chat id, and optionally yours so you can check in without asking him).
ALLOWED_CHAT_IDS = {
    int(cid.strip()) for cid in os.environ["ALLOWED_CHAT_IDS"].split(",") if cid.strip()
}
# A human label for which job site this instance is, shown in replies — e.g. "Site 1 - Main St".
SITE_NAME = os.environ.get("SITE_NAME", "Site")

API_URL = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}"
XLSX_PATH = os.environ.get("XLSX_LOCAL_PATH", "expenses.xlsx")

# --- Excel styling ---
HEADER_FILL = PatternFill(start_color="2F6F5E", end_color="2F6F5E", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill(start_color="DCEAE5", end_color="DCEAE5", fill_type="solid")
THIN_SIDE = Side(style="thin", color="B7B7B7")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

pending_confirmations = {}

HELP_TEXT = (
    f"Hey! This is the *{SITE_NAME}* expense bot. Just text me an amount whenever you spend "
    "something on this site, e.g.:\n\n"
    "  250 lumber for framing\n"
    "  eighty bucks gas for the truck\n"
    "  1200 paid the electrician\n\n"
    "I'll figure out the category automatically (materials, labor, permits, etc.) — "
    "you can always correct it later by editing the Excel file and sending it back to me.\n\n"
    "Commands:\n"
    "/total - running total for this site\n"
    "/thismonth - total for the current month\n"
    "/bycategory - breakdown by category\n"
    "/download - get the Excel file\n"
    "/undo - remove the last entry\n\n"
    "Edited the Excel file yourself? Send it back to me as a file attachment and "
    "I'll use your edited version going forward.\n"
)

AMOUNT_RE = re.compile(r"[-+]?\d+(?:\.\d+)?")

NUMBER_WORDS = {
    "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
    "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
    "thirteen": 13, "fourteen": 14, "fifteen": 15, "sixteen": 16, "seventeen": 17,
    "eighteen": 18, "nineteen": 19, "twenty": 20, "thirty": 30, "forty": 40,
    "fifty": 50, "sixty": 60, "seventy": 70, "eighty": 80, "ninety": 90,
}
SCALE_WORDS = {"hundred": 100, "thousand": 1000, "million": 1000000, "billion": 1000000000}
FILLER_AFTER_AMOUNT = {"dollars", "dollar", "usd", "bucks"}


def words_to_number(words):
    total = 0
    current = 0
    for w in words:
        if w in NUMBER_WORDS:
            current += NUMBER_WORDS[w]
        elif w in SCALE_WORDS:
            scale = SCALE_WORDS[w]
            current = (current or 1) * scale
            if scale >= 1000:
                total += current
                current = 0
    return total + current


def clean_word(w):
    return re.sub(r"[^a-zA-Z]", "", w).lower()


def extract_amount_and_note(text):
    # 1) Prefer plain digits, e.g. "250 lumber for framing"
    match = AMOUNT_RE.search(text)
    if match:
        amount = float(match.group())
        note = (text[:match.start()] + text[match.end():]).strip(" $-")
        return amount, note

    # 2) Fall back to spelled-out numbers, e.g. "eighty bucks gas for the truck"
    words = text.split()
    n = len(words)
    best_start = best_end = None
    i = 0
    while i < n:
        if clean_word(words[i]) in NUMBER_WORDS or clean_word(words[i]) in SCALE_WORDS:
            start = i
            j = i
            while j < n and (clean_word(words[j]) in NUMBER_WORDS or clean_word(words[j]) in SCALE_WORDS):
                j += 1
            if best_start is None or (j - start) > (best_end - best_start):
                best_start, best_end = start, j
            i = j
        else:
            i += 1

    if best_start is None:
        return None, text

    number_words = [clean_word(w) for w in words[best_start:best_end]]
    amount = words_to_number(number_words)
    if amount <= 0:
        return None, text

    remaining = words[:best_start] + words[best_end:]
    if remaining and clean_word(remaining[0]) in FILLER_AFTER_AMOUNT:
        remaining = remaining[1:]
    note = " ".join(remaining).strip(" $-")
    return amount, note


def is_total_row(values):
    return bool(values) and values[0] is not None and str(values[0]).strip().upper() == "TOTAL"


def get_last_entry(ws):
    for r in range(ws.max_row, 1, -1):
        values = [c.value for c in ws[r]]
        if is_total_row(values):
            continue
        return values
    return None


def is_recent_duplicate(amount, minutes=10):
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb["Log"]
    last = get_last_entry(ws)
    if not last or last[1] is None:
        return False
    try:
        last_dt = datetime.strptime(str(last[0]), "%Y-%m-%d %H:%M")
    except ValueError:
        return False
    same_amount = abs(float(last[1]) - float(amount)) < 0.01
    recent = (datetime.now() - last_dt).total_seconds() <= minutes * 60
    return same_amount and recent


def get_category_breakdown():
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb["Log"]
    totals = defaultdict(float)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if is_total_row(row):
            continue
        amount, category = row[1], row[2]
        if amount is None:
            continue
        totals[category or "Other"] += float(amount)
    return totals


def get_month_total(year, month):
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb["Log"]
    total, count = 0.0, 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        date_str, amount = row[0], row[1]
        if not date_str or amount is None:
            continue
        try:
            dt = datetime.strptime(str(date_str), "%Y-%m-%d %H:%M")
        except ValueError:
            continue
        if dt.year == year and dt.month == month:
            total += float(amount)
            count += 1
    return total, count


def style_header(ws, ncols=4):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def apply_borders(ws, ncols=4):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ncols):
        if all(c.value is None for c in row):
            continue
        for c in row:
            c.border = CELL_BORDER


def strip_total_row(ws):
    if ws.max_row > 1:
        values = [c.value for c in ws[ws.max_row]]
        if is_total_row(values):
            ws.delete_rows(ws.max_row)


def add_log_total_row(ws):
    total = 0.0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[1] is not None:
            total += float(row[1])
    ws.append(["TOTAL", round(total, 2), "", ""])
    r = ws.max_row
    for col in range(1, 5):
        c = ws.cell(row=r, column=col)
        c.font = Font(bold=True)
        c.fill = TOTAL_FILL


def finalize_log_sheet(ws):
    strip_total_row(ws)
    if ws.max_row > 1:
        add_log_total_row(ws)
    style_header(ws)
    apply_borders(ws)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 45


def ensure_workbook():
    if not os.path.exists(XLSX_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Log"
        ws.append(["Date", "Amount", "Category", "Note"])
        style_header(ws)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 45
        wb.save(XLSX_PATH)


def rebuild_summary(wb):
    log_ws = wb["Log"]
    monthly = defaultdict(lambda: defaultdict(float))

    for row in log_ws.iter_rows(min_row=2, values_only=True):
        if is_total_row(row):
            continue
        date_str, amount = row[0], row[1]
        if not date_str or amount is None:
            continue
        try:
            dt = datetime.strptime(str(date_str), "%Y-%m-%d %H:%M")
        except ValueError:
            continue
        monthly[dt.year][dt.month] += float(amount)

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary")
    ws.append(["Year", "Month", "Total"])
    style_header(ws, ncols=3)

    for year in sorted(monthly.keys()):
        year_total = 0.0
        for month in sorted(monthly[year].keys()):
            total = monthly[year][month]
            year_total += total
            ws.append([year, MONTH_NAMES[month - 1], round(total, 2)])
        total_row = ws.max_row + 1
        ws.append([year, "Year Total", round(year_total, 2)])
        for col in ("A", "B", "C"):
            ws[f"{col}{total_row}"].font = Font(bold=True)
            ws[f"{col}{total_row}"].fill = TOTAL_FILL
        ws.append([])

    apply_borders(ws, ncols=3)
    for col, width in zip("ABC", (8, 14, 14)):
        ws.column_dimensions[col].width = width


def append_entry(amount, category, note):
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb["Log"]
    strip_total_row(ws)
    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M"), amount, category, note])
    finalize_log_sheet(ws)
    rebuild_summary(wb)
    wb.save(XLSX_PATH)
    return github_store.push_current_with_retry(XLSX_PATH, f"Log ${amount:,.2f} ({category}) — {SITE_NAME}")


def undo_last():
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb["Log"]
    strip_total_row(ws)
    if ws.max_row <= 1:
        return None
    last_row = [c.value for c in ws[ws.max_row]]
    ws.delete_rows(ws.max_row)
    finalize_log_sheet(ws)
    rebuild_summary(wb)
    wb.save(XLSX_PATH)
    github_store.push_current_with_retry(XLSX_PATH, f"Undo last entry — {SITE_NAME}")
    return last_row


def get_total():
    ensure_workbook()
    wb = load_workbook(XLSX_PATH)
    ws = wb["Log"]
    total = 0.0
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if is_total_row(row):
            continue
        if row[1] is not None:
            total += float(row[1])
            count += 1
    return total, count


def send_message(chat_id, text):
    requests.post(f"{API_URL}/sendMessage", json={"chat_id": chat_id, "text": text})


def send_document(chat_id):
    ensure_workbook()
    with open(XLSX_PATH, "rb") as f:
        requests.post(
            f"{API_URL}/sendDocument",
            data={"chat_id": chat_id},
            files={"document": (f"{SITE_NAME.replace(' ', '_')}_expenses.xlsx", f)},
        )


def handle_incoming_document(chat_id, document):
    file_name = document.get("file_name", "")
    if not file_name.lower().endswith(".xlsx"):
        send_message(chat_id, "That doesn't look like an .xlsx file, so I left your data untouched.")
        return

    file_id = document.get("file_id")
    info = requests.get(f"{API_URL}/getFile", params={"file_id": file_id}).json()
    if not info.get("ok"):
        send_message(chat_id, "Couldn't fetch that file from Telegram, try sending it again.")
        return

    file_path = info["result"]["file_path"]
    file_url = f"https://api.telegram.org/file/bot{TELEGRAM_TOKEN}/{file_path}"
    resp = requests.get(file_url)

    tmp_path = "incoming_tmp.xlsx"
    with open(tmp_path, "wb") as f:
        f.write(resp.content)

    try:
        wb = load_workbook(tmp_path)
        if "Log" not in wb.sheetnames:
            raise ValueError("Missing Log sheet")
    except Exception:
        os.remove(tmp_path)
        send_message(chat_id, "I couldn't read that file properly (it needs a 'Log' sheet), so I kept your old data.")
        return

    ws = wb["Log"]
    finalize_log_sheet(ws)
    rebuild_summary(wb)
    wb.save(tmp_path)
    os.replace(tmp_path, XLSX_PATH)
    ok = github_store.push_current_with_retry(XLSX_PATH, f"Manual edit uploaded via Telegram — {SITE_NAME}")
    total, count = get_total()
    note = "" if ok or not github_store.ENABLED else " (heads up: couldn't save this to GitHub, will retry on next entry)"
    send_message(chat_id, f"Got it — saved your edits. Running total is now ${total:,.2f} ({count} entries).{note} New entries will build on top of this.")


@app.route("/webhook", methods=["POST"])
def webhook():
    update = request.get_json(force=True, silent=True) or {}
    message = update.get("message", {})
    chat_id = message.get("chat", {}).get("id")

    if chat_id is None:
        return "ok"

    if chat_id not in ALLOWED_CHAT_IDS:
        return "ok"

    if "document" in message:
        handle_incoming_document(chat_id, message["document"])
        return "ok"

    text = (message.get("text") or "").strip()
    if not text:
        return "ok"

    if chat_id in pending_confirmations:
        if text.lower() in ("yes", "y"):
            pending_amount, pending_category, pending_note = pending_confirmations.pop(chat_id)
            append_entry(pending_amount, pending_category, pending_note)
            total, count = get_total()
            send_message(chat_id, f"Logged ${pending_amount:,.2f} ({pending_category}). Running total: ${total:,.2f} ({count} entries).")
            return "ok"
        elif text.lower() in ("no", "n", "cancel"):
            pending_confirmations.pop(chat_id, None)
            send_message(chat_id, "Cancelled, didn't log that one.")
            return "ok"
        else:
            pending_confirmations.pop(chat_id, None)
            # fall through, treat this as a fresh message

    if text in ("/start", "/help"):
        send_message(chat_id, HELP_TEXT)
        return "ok"

    if text == "/total":
        total, count = get_total()
        send_message(chat_id, f"{SITE_NAME} total: ${total:,.2f} across {count} entries.")
        return "ok"

    if text == "/thismonth":
        now = datetime.now()
        total, count = get_month_total(now.year, now.month)
        send_message(chat_id, f"{SITE_NAME} — {MONTH_NAMES[now.month - 1]} {now.year}: ${total:,.2f} across {count} entries.")
        return "ok"

    if text == "/bycategory":
        totals = get_category_breakdown()
        if not totals:
            send_message(chat_id, "No entries yet.")
            return "ok"
        lines = [f"{cat}: ${amt:,.2f}" for cat, amt in sorted(totals.items(), key=lambda x: -x[1])]
        send_message(chat_id, f"{SITE_NAME} — totals by category:\n" + "\n".join(lines))
        return "ok"

    if text == "/download":
        send_document(chat_id)
        return "ok"

    if text == "/undo":
        removed = undo_last()
        if removed:
            send_message(chat_id, f"Removed: {removed[0]} — ${removed[1]} ({removed[2]}) {removed[3] or 'no note'}")
        else:
            send_message(chat_id, "Nothing to undo.")
        return "ok"

    if text == "/categories":
        send_message(chat_id, "Categories I recognize:\n" + "\n".join(f"- {c}" for c in category_list()))
        return "ok"

    amount, note = extract_amount_and_note(text)
    if amount is None:
        send_message(chat_id, "Didn't catch an amount there. Try something like: 250 lumber for framing or eighty bucks gas")
        return "ok"

    category = detect_category(note)

    if is_recent_duplicate(amount):
        pending_confirmations[chat_id] = (amount, category, note)
        send_message(chat_id, f"Heads up — you just logged ${amount:,.2f} in the last few minutes too. Log this one as well? Reply yes or no.")
        return "ok"

    ok = append_entry(amount, category, note)
    total, count = get_total()
    warn = "" if ok or not github_store.ENABLED else " ⚠️ couldn't back this up to GitHub yet, will retry"
    send_message(chat_id, f"Logged ${amount:,.2f} as *{category}*. Running total: ${total:,.2f} ({count} entries).{warn}")
    return "ok"


@app.route("/", methods=["GET"])
def health():
    return f"{SITE_NAME} expense bot is running."


def _bootstrap():
    """Runs once when the process starts: pull the real data from GitHub before
    anything else touches the local (ephemeral) file."""
    try:
        found = github_store.pull_latest(XLSX_PATH)
        if not found:
            ensure_workbook()
    except Exception:
        # If GitHub is briefly unreachable at boot, don't crash — fall back to
        # whatever's on local disk (or a fresh workbook) and keep trying to push on writes.
        ensure_workbook()


_bootstrap()


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
